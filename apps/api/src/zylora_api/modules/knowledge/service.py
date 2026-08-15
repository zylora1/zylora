from __future__ import annotations

import hashlib
import io
import json
import logging
import re
import socket
import struct
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime, timedelta
from pathlib import PurePath
from typing import Protocol
from uuid import UUID, uuid4

from docx import Document
from pypdf import PdfReader
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession
from zylora_api.core.config import Settings
from zylora_api.db.knowledge_models import KnowledgeSource
from zylora_api.db.models import OutboxEvent
from zylora_api.db.website_models import Website
from zylora_api.modules.templates.service import problem
from zylora_api.storage.base import ObjectStorage

SAFE_NAME = re.compile(r"[^A-Za-z0-9._ -]+")
logger = logging.getLogger("zylora.knowledge")

CONTROL_CHARACTERS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
MAX_ATTEMPTS = 4

ALLOWED_DOCUMENTS = {
    ".pdf": ("PDF", {"application/pdf"}),
    ".docx": (
        "DOCX",
        {
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/octet-stream",
        },
    ),
    ".txt": ("TEXT", {"text/plain"}),
    ".md": ("MARKDOWN", {"text/markdown", "text/plain"}),
}


class DocumentSafetyError(ValueError):
    def __init__(self, code: str, safe_message: str) -> None:
        super().__init__(safe_message)
        self.code = code
        self.safe_message = safe_message


class RetryableScanError(RuntimeError):
    pass


class DocumentScanner(Protocol):
    def scan(self, data: bytes) -> None: ...


class TestDocumentScanner:
    def scan(self, data: bytes) -> None:
        if b"EICAR-STANDARD-ANTIVIRUS-TEST-FILE" in data:
            raise DocumentSafetyError(
                "malware_detected", "The document did not pass the safety scan."
            )


class ClamAVDocumentScanner:
    """Small INSTREAM adapter; scanner-specific details stay behind this boundary."""

    def __init__(self, host: str, port: int, timeout: float) -> None:
        self.host = host
        self.port = port
        self.timeout = timeout

    def scan(self, data: bytes) -> None:
        try:
            with socket.create_connection((self.host, self.port), timeout=self.timeout) as client:
                client.sendall(b"zINSTREAM\0")
                for offset in range(0, len(data), 64 * 1024):
                    chunk = data[offset : offset + 64 * 1024]
                    client.sendall(struct.pack(">I", len(chunk)))
                    client.sendall(chunk)
                client.sendall(struct.pack(">I", 0))
                response = client.recv(4096).decode("utf-8", errors="replace")
        except (OSError, TimeoutError) as error:
            raise RetryableScanError("document scanner unavailable") from error
        if "FOUND" in response:
            raise DocumentSafetyError(
                "malware_detected", "The document did not pass the safety scan."
            )
        if "OK" not in response:
            raise RetryableScanError("document scanner returned an invalid response")


@dataclass(frozen=True)
class ValidatedDocument:
    source_type: str
    mime_type: str
    safe_name: str
    checksum: str


@dataclass(frozen=True)
class ExtractedSection:
    text: str
    page_number: int | None = None
    heading: str | None = None


def scanner_for(settings: Settings) -> DocumentScanner:
    if settings.document_scanner_provider == "test":
        return TestDocumentScanner()
    return ClamAVDocumentScanner(
        settings.clamav_host, settings.clamav_port, settings.clamav_timeout_seconds
    )


def _safe_filename(filename: str) -> str:
    basename = PurePath(filename.replace("\\", "/")).name
    safe = SAFE_NAME.sub("_", basename).strip(" .")
    if not safe:
        safe = "document"
    stem, dot, suffix = safe.rpartition(".")
    if dot and len(stem) > 180:
        safe = f"{stem[:180]}.{suffix}"
    return safe[:240]


def validate_document(
    *, filename: str, content_type: str | None, data: bytes, settings: Settings
) -> ValidatedDocument:
    if not data:
        raise DocumentSafetyError("document_empty", "The document is empty.")
    if len(data) > settings.knowledge_max_file_bytes:
        raise DocumentSafetyError("document_too_large", "The document is too large.")
    safe_name = _safe_filename(filename)
    extension = PurePath(safe_name).suffix.casefold()
    definition = ALLOWED_DOCUMENTS.get(extension)
    if not definition:
        if extension == ".doc":
            raise DocumentSafetyError(
                "legacy_doc_unsupported", "Upload a DOCX or PDF instead of a legacy DOC file."
            )
        raise DocumentSafetyError(
            "document_unsupported", "Upload a PDF, DOCX, TXT, or Markdown file."
        )
    source_type, accepted_mimes = definition
    normalized_mime = (content_type or "").split(";", 1)[0].strip().casefold()
    if normalized_mime not in accepted_mimes:
        raise DocumentSafetyError(
            "document_type_mismatch", "The file extension and content type do not match."
        )
    if extension == ".pdf" and not data.startswith(b"%PDF-"):
        raise DocumentSafetyError("document_type_mismatch", "This is not a valid PDF file.")
    if extension == ".docx":
        if not data.startswith(b"PK\x03\x04"):
            raise DocumentSafetyError("document_type_mismatch", "This is not a valid DOCX file.")
        _inspect_docx_archive(data, settings)
    if extension in {".txt", ".md"}:
        _decode_text(data, settings)
    return ValidatedDocument(
        source_type=source_type,
        mime_type={
            "PDF": "application/pdf",
            "DOCX": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "TEXT": "text/plain",
            "MARKDOWN": "text/markdown",
        }[source_type],
        safe_name=safe_name,
        checksum=hashlib.sha256(data).hexdigest(),
    )


def _inspect_docx_archive(data: bytes, settings: Settings) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > settings.knowledge_max_docx_entries:
                raise DocumentSafetyError("document_archive_limit", "The DOCX is too complex.")
            total = sum(item.file_size for item in entries)
            if any(item.flag_bits & 0x1 for item in entries):
                raise DocumentSafetyError(
                    "document_encrypted",
                    "Encrypted DOCX files are not supported.",
                )
            if total > settings.knowledge_max_docx_uncompressed_bytes:
                raise DocumentSafetyError(
                    "document_archive_limit", "The DOCX expands beyond the safety limit."
                )
            names = {item.filename.casefold() for item in entries}
            if "[content_types].xml" not in names or "word/document.xml" not in names:
                raise DocumentSafetyError("document_malformed", "The DOCX file is malformed.")
            blocked = ("vbaproject.bin", "activex/", "embeddings/")
            if any(any(marker in name for marker in blocked) for name in names):
                raise DocumentSafetyError(
                    "document_active_content",
                    "DOCX files with embedded active content are not supported.",
                )
            if any(
                item.filename.startswith(("/", "\\")) or ".." in PurePath(item.filename).parts
                for item in entries
            ):
                raise DocumentSafetyError("document_malformed", "The DOCX file is malformed.")
    except zipfile.BadZipFile as error:
        raise DocumentSafetyError("document_malformed", "The DOCX file is malformed.") from error


def _normalize_text(value: str) -> str:
    value = CONTROL_CHARACTERS.sub("", value.replace("\r\n", "\n").replace("\r", "\n"))
    lines = [" ".join(line.split()) for line in value.split("\n")]
    output: list[str] = []
    for line in lines:
        if line or (output and output[-1]):
            output.append(line)
    return "\n".join(output).strip()


def _decode_text(data: bytes, settings: Settings) -> str:
    try:
        value = data.decode("utf-8-sig", errors="strict")
    except UnicodeDecodeError as error:
        raise DocumentSafetyError(
            "document_encoding", "Text documents must use UTF-8 encoding."
        ) from error
    value = _normalize_text(value)
    if not value:
        raise DocumentSafetyError("document_no_text", "No machine-readable text was found.")
    if len(value) > settings.knowledge_max_extracted_characters:
        raise DocumentSafetyError("document_text_limit", "The document contains too much text.")
    return value


def extract_document(source_type: str, data: bytes, settings: Settings) -> list[ExtractedSection]:
    sections: list[ExtractedSection]
    if source_type in {"TEXT", "MARKDOWN"}:
        return [ExtractedSection(_decode_text(data, settings))]
    if source_type == "PDF":
        try:
            reader = PdfReader(io.BytesIO(data), strict=True)
            if len(reader.pages) > settings.knowledge_max_pdf_pages:
                raise DocumentSafetyError("document_page_limit", "The PDF has too many pages.")
            sections = [
                ExtractedSection(_normalize_text(page.extract_text() or ""), page_number=index + 1)
                for index, page in enumerate(reader.pages)
            ]
        except DocumentSafetyError:
            raise
        except Exception as error:
            raise DocumentSafetyError("document_malformed", "The PDF file is malformed.") from error
        sections = [item for item in sections if item.text]
        if not sections:
            raise DocumentSafetyError(
                "document_no_text", "No machine-readable text was found in this PDF."
            )
    elif source_type == "DOCX":
        _inspect_docx_archive(data, settings)
        try:
            document = Document(io.BytesIO(data))
            sections = []
            heading: str | None = None
            for paragraph in document.paragraphs:
                text = _normalize_text(paragraph.text)
                if not text:
                    continue
                if paragraph.style and paragraph.style.name.casefold().startswith("heading"):
                    heading = text
                sections.append(ExtractedSection(text=text, heading=heading))
            for table in document.tables:
                rows = [
                    " | ".join(_normalize_text(cell.text) for cell in row.cells)
                    for row in table.rows
                ]
                table_text = _normalize_text("\n".join(rows))
                if table_text:
                    sections.append(ExtractedSection(text=table_text, heading=heading))
        except Exception as error:
            raise DocumentSafetyError(
                "document_malformed", "The DOCX file is malformed."
            ) from error
        if not sections:
            raise DocumentSafetyError("document_no_text", "No machine-readable text was found.")
    else:
        raise DocumentSafetyError("document_unsupported", "This document type is unsupported.")
    if sum(len(item.text) for item in sections) > settings.knowledge_max_extracted_characters:
        raise DocumentSafetyError("document_text_limit", "The document contains too much text.")
    return sections


class KnowledgeSourceService:
    def __init__(
        self,
        session: AsyncSession,
        storage: ObjectStorage,
        settings: Settings,
        scanner: DocumentScanner | None = None,
    ) -> None:
        self.session = session
        self.storage = storage
        self.settings = settings
        self.scanner = scanner or scanner_for(settings)

    async def list_for_owner(self, website_id: UUID, owner_user_id: UUID) -> list[KnowledgeSource]:
        await self._owned_website(website_id, owner_user_id)
        return list(
            (
                await self.session.scalars(
                    select(KnowledgeSource)
                    .where(
                        KnowledgeSource.website_id == website_id,
                        KnowledgeSource.owner_user_id == owner_user_id,
                        KnowledgeSource.status != "DELETED",
                    )
                    .order_by(KnowledgeSource.created_at.desc())
                )
            ).all()
        )

    async def create(
        self,
        *,
        website_id: UUID,
        owner_user_id: UUID,
        filename: str,
        content_type: str | None,
        data: bytes,
        correlation_id: str,
    ) -> KnowledgeSource:
        if not self.settings.knowledge_ingestion_enabled:
            raise problem(
                503, "knowledge_ingestion_disabled", "Document ingestion is not configured."
            )
        website = await self._owned_website(website_id, owner_user_id, lock=True)
        count = int(
            await self.session.scalar(
                select(func.count(KnowledgeSource.id)).where(
                    KnowledgeSource.website_id == website.id,
                    KnowledgeSource.status != "DELETED",
                )
            )
            or 0
        )
        if count >= self.settings.knowledge_max_documents_per_website:
            raise problem(
                422, "knowledge_document_limit", "This Website has reached its document limit."
            )
        try:
            validated = validate_document(
                filename=filename, content_type=content_type, data=data, settings=self.settings
            )
        except DocumentSafetyError as error:
            raise problem(422, error.code, error.safe_message) from error
        source_id = uuid4()
        key = (
            f"knowledge/{website.id}/{owner_user_id}/{source_id}/"
            f"{validated.checksum}/{validated.safe_name}"
        )
        metadata = self.storage.put_bytes(key, data, validated.mime_type)
        source = KnowledgeSource(
            id=source_id,
            website_id=website.id,
            owner_user_id=owner_user_id,
            source_type=validated.source_type,
            original_filename=filename[:500],
            safe_display_name=validated.safe_name,
            storage_key=metadata.key,
            mime_type=validated.mime_type,
            byte_size=metadata.size,
            sha256=metadata.checksum_sha256,
            status="UPLOADED",
        )
        self.session.add(source)
        self.session.add(
            OutboxEvent(
                aggregate_type="KNOWLEDGE_SOURCE",
                aggregate_id=source.id,
                event_type="knowledge.source_ingest_requested",
                payload={"knowledge_source_id": str(source.id)},
                correlation_id=correlation_id,
            )
        )
        logger.info(
            "knowledge_uploaded",
            extra={
                "website_id": str(website.id),
                "source_id": str(source.id),
                "event_type": "knowledge.uploaded",
                "correlation_id": correlation_id,
                "artifact_bytes": source.byte_size,
            },
        )
        return source

    async def retry(
        self, source_id: UUID, website_id: UUID, owner_user_id: UUID, correlation_id: str
    ) -> KnowledgeSource:
        source = await self._owned_source(source_id, website_id, owner_user_id, lock=True)
        if source.status != "FAILED":
            raise problem(409, "knowledge_retry_invalid", "Only failed documents can be retried.")
        source.status = "UPLOADED"
        source.failure_code = None
        source.failure_message_safe = None
        source.version += 1
        self.session.add(
            OutboxEvent(
                aggregate_type="KNOWLEDGE_SOURCE",
                aggregate_id=source.id,
                event_type="knowledge.source_ingest_requested",
                payload={"knowledge_source_id": str(source.id)},
                correlation_id=correlation_id,
            )
        )
        return source

    async def delete(
        self, source_id: UUID, website_id: UUID, owner_user_id: UUID, correlation_id: str
    ) -> KnowledgeSource:
        source = await self._owned_source(source_id, website_id, owner_user_id, lock=True)
        if source.status == "DELETED":
            return source
        source.status = "DELETED"
        source.deleted_at = datetime.now(UTC)
        source.version += 1
        self.session.add(
            OutboxEvent(
                aggregate_type="KNOWLEDGE_SOURCE",
                aggregate_id=source.id,
                event_type="knowledge.source_delete_requested",
                payload={"knowledge_source_id": str(source.id)},
                correlation_id=correlation_id,
            )
        )
        return source

    async def process_outbox_event(self, event_id: UUID) -> OutboxEvent:
        event = await self.session.scalar(
            select(OutboxEvent).where(OutboxEvent.id == event_id).with_for_update()
        )
        if not event:
            raise problem(404, "knowledge_event_not_found", "Knowledge job not found.")
        if event.state == "PUBLISHED":
            return event
        event.attempts += 1
        try:
            source = await self.session.scalar(
                select(KnowledgeSource)
                .where(KnowledgeSource.id == UUID(str(event.payload["knowledge_source_id"])))
                .with_for_update()
            )
            if not source:
                raise DocumentSafetyError("knowledge_source_missing", "Knowledge source not found.")
            if event.event_type == "knowledge.source_delete_requested":
                self.storage.delete(source.storage_key)
                if source.extracted_storage_key:
                    self.storage.delete(source.extracted_storage_key)
                await self._request_rebuild(source, event.correlation_id)
            elif event.event_type == "knowledge.source_ingest_requested":
                if source.status == "DELETED":
                    event.state = "PUBLISHED"
                    return event
                source.status = "SCANNING"
                data = self.storage.get_bytes(source.storage_key)
                if hashlib.sha256(data).hexdigest() != source.sha256:
                    raise DocumentSafetyError(
                        "document_checksum_mismatch", "The stored document is invalid."
                    )
                self.scanner.scan(data)
                source.status = "PROCESSING"
                sections = extract_document(source.source_type, data, self.settings)
                payload = {
                    "source_id": str(source.id),
                    "source_version": source.version,
                    "sections": [
                        {
                            "text": item.text,
                            "page_number": item.page_number,
                            "heading": item.heading,
                        }
                        for item in sections
                    ],
                }
                extracted = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode()
                extracted_key = (
                    f"knowledge-extracted/{source.website_id}/{source.owner_user_id}/"
                    f"{source.id}/v{source.version}.json"
                )
                self.storage.put_bytes(extracted_key, extracted, "application/json")
                source.extracted_storage_key = extracted_key
                source.extraction_metadata = {
                    "section_count": len(sections),
                    "character_count": sum(len(item.text) for item in sections),
                }
                source.status = "READY"
                logger.info(
                    "knowledge_processing_completed",
                    extra={
                        "website_id": str(source.website_id),
                        "source_id": str(source.id),
                        "event_type": "knowledge.processing_completed",
                        "correlation_id": event.correlation_id,
                        "outcome": "ready",
                    },
                )
                source.failure_code = None
                source.failure_message_safe = None
                await self._request_rebuild(source, event.correlation_id)
            else:
                raise DocumentSafetyError(
                    "knowledge_event_unsupported", "Knowledge job is unsupported."
                )
            event.state = "PUBLISHED"
            event.published_at = datetime.now(UTC)
            event.last_error_code = None
        except RetryableScanError:
            if event.attempts >= MAX_ATTEMPTS:
                await self._fail_source(
                    event, "document_scan_unavailable", "Document safety scanning is unavailable."
                )
                event.state = "FAILED"
            else:
                event.state = "PENDING"
                event.available_at = datetime.now(UTC) + timedelta(
                    minutes=2 ** (event.attempts - 1)
                )
                event.last_error_code = "document_scan_retry_wait"
        except (DocumentSafetyError, OSError, RuntimeError, KeyError, ValueError) as error:
            code = (
                error.code
                if isinstance(error, DocumentSafetyError)
                else "document_processing_failed"
            )
            message = (
                error.safe_message
                if isinstance(error, DocumentSafetyError)
                else "Document could not be processed."
            )
            await self._fail_source(event, code, message)
            event.state = "FAILED"
            logger.warning(
                "knowledge_processing_failed",
                extra={
                    "source_id": str(event.payload.get("knowledge_source_id", "")),
                    "event_type": "knowledge.processing_failed",
                    "correlation_id": event.correlation_id,
                    "attempt": event.attempts,
                    "outcome": "failed",
                    "safe_error_code": code,
                },
            )
        finally:
            event.lease_owner = None
            event.leased_until = None
        return event

    async def _request_rebuild(self, source: KnowledgeSource, correlation_id: str) -> None:
        website = await self.session.get(Website, source.website_id)
        if not website or website.status != "PUBLISHED" or not website.published_version_id:
            return
        from zylora_api.modules.chatbot.indexing import KnowledgeIndexService

        await KnowledgeIndexService(self.session, self.storage, self.settings).request_rebuild(
            website, website.published_version_id, correlation_id
        )

    async def _fail_source(self, event: OutboxEvent, code: str, message: str) -> None:
        source_id = event.payload.get("knowledge_source_id")
        if source_id:
            source = await self.session.get(KnowledgeSource, UUID(str(source_id)))
            if source and source.status != "DELETED":
                source.status = "FAILED"
                source.failure_code = code
                source.failure_message_safe = message
        event.last_error_code = code

    async def _owned_website(
        self, website_id: UUID, owner_user_id: UUID, lock: bool = False
    ) -> Website:
        query = select(Website).where(
            Website.id == website_id, Website.owner_user_id == owner_user_id
        )
        if lock:
            query = query.with_for_update()
        website = await self.session.scalar(query)
        if not website:
            raise problem(404, "website_not_found", "Website not found.")
        return website

    async def _owned_source(
        self, source_id: UUID, website_id: UUID, owner_user_id: UUID, lock: bool
    ) -> KnowledgeSource:
        query = select(KnowledgeSource).where(
            KnowledgeSource.id == source_id,
            KnowledgeSource.website_id == website_id,
            KnowledgeSource.owner_user_id == owner_user_id,
        )
        if lock:
            query = query.with_for_update()
        source = await self.session.scalar(query)
        if not source:
            raise problem(404, "knowledge_source_not_found", "Knowledge source not found.")
        return source
